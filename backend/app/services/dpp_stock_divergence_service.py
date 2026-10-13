"""Investigação 4.1: saldos por depósito fora do consolidado do DPP.

Este módulo é somente leitura. Quantidades nos depósitos investigados não
constituem estoque disponível até confirmação do PCM com Almoxarifado.
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from app.services.dpp_consolidation_service import _material_key, _normalize, _number

SOURCE_SHEET = "ORIGINAL"
TARGET_DEPOTS = ("M221", "P277", "G702")
REQUIRED_COLUMNS = {
    "material": "Material",
    "depot": "Depósito",
    "quantity": "Utilização livre",
    "unit": "Unid.medida básica",
}
OPTIONAL_COLUMNS = {"center": "Centro"}


def _column_indexes(header_row: tuple) -> dict[str, int | None]:
    normalized = {_normalize(value): index for index, value in enumerate(header_row) if value is not None}
    columns: dict[str, int | None] = {}
    for field, label in REQUIRED_COLUMNS.items():
        index = normalized.get(_normalize(label))
        if index is None:
            raise ValueError(f"A aba ORIGINAL não contém a coluna obrigatória '{label}'.")
        columns[field] = index
    for field, label in OPTIONAL_COLUMNS.items():
        columns[field] = normalized.get(_normalize(label))
    return columns


def _cell(row: tuple, index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def investigate_stock_divergence(
    *,
    stock_content: bytes,
    stock_filename: str,
    scenario: dict,
) -> dict:
    """Cruza PNs críticos do cenário com M221/P277/G702 na aba ORIGINAL.

    Não altera cenário, estoque, NEC, SALDO, REAL nem gap. O resultado é uma
    evidência preliminar para validação física/humana, não uma liberação.
    """
    if Path(stock_filename or "").suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Envie o mesmo arquivo STK SAP (.xlsx ou .xlsm) usado no cenário.")
    if not stock_content:
        raise ValueError("O arquivo STK SAP está vazio.")

    scenario_sources = scenario.get("sources") or []
    stock_source = next((item for item in scenario_sources if item.get("id") == "stock"), None)
    expected_filename = (stock_source or {}).get("filename")
    if expected_filename and expected_filename.casefold() != stock_filename.casefold():
        raise ValueError("O STK enviado não corresponde ao nome do arquivo usado para gerar este cenário.")
    expected_fingerprint = (stock_source or {}).get("fingerprint_sha256")
    if expected_fingerprint and sha256(stock_content).hexdigest() != expected_fingerprint:
        raise ValueError("O conteúdo do STK difere do arquivo usado na geração do cenário. Refaça a geração com o pacote correto.")

    critical = {
        key: material
        for material in scenario.get("materials") or []
        if material.get("status") == "INVESTIGAR"
        if (key := _material_key(material.get("material") or material.get("material_key")))
    }
    counts = {depot: 0 for depot in TARGET_DEPOTS}
    grouped: dict[tuple[str, str, str, str], dict] = {}
    ignored = defaultdict(int)

    try:
        workbook = load_workbook(BytesIO(stock_content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Não foi possível abrir a planilha STK SAP.") from exc

    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError("A aba 'ORIGINAL' não existe no STK SAP. A geração normal do DPP não foi afetada.")

        sheet = workbook[SOURCE_SHEET]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise ValueError("A aba 'ORIGINAL' do STK SAP está vazia.")
        columns = _column_indexes(header_row)
        for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            depot = str(_cell(row, columns["depot"]) or "").strip().upper()
            if depot not in counts:
                continue
            counts[depot] += 1

            material_key = _material_key(_cell(row, columns["material"]))
            if material_key not in critical:
                continue

            quantity_raw = _cell(row, columns["quantity"])
            quantity = _number(quantity_raw, None)
            if quantity is None or isinstance(quantity_raw, bool):
                ignored["quantidade_invalida"] += 1
                continue
            if quantity <= 0:
                ignored["quantidade_nao_positiva"] += 1
                continue

            unit = str(_cell(row, columns["unit"]) or "").strip().upper()
            expected_unit = str(critical[material_key].get("um") or "").strip().upper()
            if not unit or unit != expected_unit:
                ignored["unidade_divergente"] += 1
                continue

            center = str(_cell(row, columns["center"]) or "").strip()
            group_key = (material_key, depot, center, unit)
            if group_key not in grouped:
                grouped[group_key] = {
                    "depot": depot,
                    "center": center or None,
                    "unit": unit,
                    "quantity": 0.0,
                    "source_rows": [],
                    "validation_status": "AGUARDANDO_PCM_ALMOXARIFADO",
                }
            item = grouped[group_key]
            item["quantity"] += quantity
            item["source_rows"].append(excel_row)
    finally:
        workbook.close()

    by_material: dict[str, list[dict]] = defaultdict(list)
    for (material_key, _, _, _), item in grouped.items():
        by_material[material_key].append(item)

    items: list[dict] = []
    for material_key, candidates in by_material.items():
        material = critical[material_key]
        candidates.sort(key=lambda item: (item["depot"], item["center"] or ""))
        items.append({
            "material": material.get("material") or material_key,
            "description": material.get("description"),
            "unit": material.get("um"),
            "balance": material.get("balance"),
            "negative_balance": max(-(_number(material.get("balance"), 0.0) or 0.0), 0.0),
            "candidate_quantity": sum(item["quantity"] for item in candidates),
            "candidates": candidates,
            "status": "PENDENTE_VALIDACAO_HUMANA",
        })
    items.sort(key=lambda item: (-item["negative_balance"], str(item["material"])))

    return {
        "scenario_id": scenario.get("scenario_id"),
        "reference_month": scenario.get("reference_month"),
        "source": {
            "filename": stock_filename,
            "sheet": SOURCE_SHEET,
            "material_column": REQUIRED_COLUMNS["material"],
            "depot_column": REQUIRED_COLUMNS["depot"],
            "quantity_column": REQUIRED_COLUMNS["quantity"],
        },
        "summary": {
            "critical_materials": len(critical),
            "materials_with_candidates": len(items),
            "materials_without_candidates": len(critical) - len(items),
            "records_by_depot": counts,
            "depots_without_records": [depot for depot, count in counts.items() if count == 0],
            "ignored_records": dict(ignored),
        },
        "materials": items,
        "requires_human_validation": True,
        "changes_dpp": False,
        "note": (
            "Quantidades brutas encontradas para investigação; não representam "
            "estoque liberado. PCM deve confirmar com Almoxarifado se o material "
            "será localizado/disponibilizado ou baixado. Nenhum SALDO foi recalculado."
        ),
    }
