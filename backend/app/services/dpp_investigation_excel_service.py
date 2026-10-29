"""Readable, auditable Excel exports for ORION stock and quality investigations.

Reports are snapshots of the generated ORION scenario and human investigation
history, not the finished DPP validation workbook and not an SAP movement.
All records are recalculated by the server before export.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from math import isfinite

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DARK = "172A40"
MID = "264664"
ACCENT = "147B90"
PALE = "EDF4F8"
WHITE = "FFFFFF"
TEXT = "263849"
MUTED = "5D6E80"
GREEN = "E3F2E8"
YELLOW = "FFF3D5"
RED = "FCE9E7"
BLUE = "E4EFF8"
BORDER = "D8E2EA"
NUMFMT = '#,##0.###;[Red](#,##0.###);–'
# Identificadores técnicos não são rótulos do relatório.
STATUS_LABELS = {
    "FINISHED": "Finalizado",
    "QUALITY_4_2": "Encaminhar à Qualidade",
    "SCRAP_4_3": "Encaminhar a Scrap",
    "AWAITING_VALIDATION": "Aguardando validação de Estoque",
    "AWAITING_QUALITY": "Aguardando validação do CQ",
}
MARKER_COLORS = {
    "Finalizado": GREEN,
    "Encaminhar à Qualidade": YELLOW,
    "Encaminhar a Scrap": YELLOW,
    "Aguardando validação de Estoque": BLUE,
    "Aguardando validação do CQ": BLUE,
}


def _status_label(value):
    return STATUS_LABELS.get(value, value or "")


def _safe(value):
    """Prevent spreadsheet formula injection from uploaded cells/user comments."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, (float, int)):
        return value if isfinite(value) else ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    if text and text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    if text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def _num(value):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if isfinite(number) else None


def _source(scenario, source_id):
    source = next((item for item in scenario.get("sources") or [] if item.get("id") == source_id), None)
    return (source or {}).get("filename") or "Não identificado no cenário"


def _header(sheet, title, subtitle, columns, widths, *, note=None):
    count = len(columns)
    last = get_column_letter(count)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
    title_cell = sheet.cell(1, 1, _safe(title))
    title_cell.fill = PatternFill("solid", fgColor=DARK)
    title_cell.font = Font(name="Aptos Display", bold=True, size=17, color=WHITE)
    title_cell.alignment = Alignment(vertical="center", indent=1)
    for cells in sheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=count):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=DARK)
    sheet.row_dimensions[1].height = 38
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count)
    sub = sheet.cell(2, 1, _safe(subtitle))
    sub.font = Font(name="Aptos", size=10, color=MUTED)
    sub.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    sheet.row_dimensions[2].height = 38 if len(subtitle) > 115 else 28
    if note:
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=count)
        warning = sheet.cell(3, 1, _safe(note))
        warning.font = Font(name="Aptos", size=10, color=MID, italic=True)
        warning.fill = PatternFill("solid", fgColor=BLUE)
        warning.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        sheet.row_dimensions[3].height = 40
    header_row = 5 if note else 4
    sheet.row_dimensions[header_row].height = 32
    for index, (label, width) in enumerate(zip(columns, widths), start=1):
        cell = sheet.cell(header_row, index, label)
        cell.font = Font(name="Aptos", bold=True, size=10, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=MID)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=ACCENT))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = f"C{header_row + 1}" if count >= 3 else f"A{header_row + 1}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    # Não comprimir 14-18 colunas numa única página ilegível.
    sheet.page_setup.fitToWidth = 3 if count >= 14 else 2 if count >= 11 else 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = True
    sheet.oddFooter.center.text = "ORION · Investigação / NÃO ALTERA O DPP"
    sheet.oddFooter.right.text = "Página &P de &N"
    return header_row


def _rows(sheet, headings, widths, rows, *, title, subtitle, note=None, status_column=None):
    header_row = _header(sheet, title, subtitle, headings, widths, note=note)
    number_cols = set()
    for index, row in enumerate(rows, start=header_row + 1):
        height = 26
        for col, raw in enumerate(row, start=1):
            val = _safe(raw)
            cell = sheet.cell(index, col, val)
            cell.fill = PatternFill("solid", fgColor=WHITE if (index - header_row) % 2 else PALE)
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=BORDER))
            if isinstance(val, (float, int)):
                cell.number_format = NUMFMT
                cell.alignment = Alignment(horizontal="right", vertical="center")
                number_cols.add(col)
            if status_column and col == status_column:
                marker = str(raw or "")
                if marker in MARKER_COLORS:
                    cell.fill = PatternFill("solid", fgColor=MARKER_COLORS[marker])
                    cell.font = Font(name="Aptos", size=10, color=TEXT, bold=True)
            if isinstance(val, str) and "\n" in val:
                height = max(height, min(135, 15 * (val.count("\n") + 1)))
        sheet.row_dimensions[index].height = height
    if not rows:
        cell = sheet.cell(header_row + 1, 1, "Nenhum registro para os critérios desta investigação.")
        cell.font = Font(name="Aptos", italic=True, color=MUTED)
        sheet.row_dimensions[header_row + 1].height = 28
    bottom = max(sheet.max_row, header_row + 1)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headings))}{bottom}"
    sheet.print_options.horizontalCentered = True


# Cada arquivo exportado espelha exatamente os cinco indicadores do respectivo
# painel. A pertinência do PN ao indicador vem do mesmo relatório usado na web.
# Um mesmo PN pode constar em várias abas sem duplicar estoque ou justificar o gap.
DECISION_LABELS = {
    "AVAILABLE": "Disponível",
    "UNAVAILABLE": "Indisponível",
    "RELEASED": "Liberado pelo CQ",
    "REJECTED": "Reprovado pelo CQ",
    "SCRAP_CONFIRMED": "Refugo confirmado",
    "NO_EVIDENCE": "Sem confirmação",
    "UNKNOWN": "A confirmar",
    "INCLUDED": "Já contabilizado no DPP",
    "NOT_INCLUDED": "Não contabilizado no DPP",
}

STOCK_COLUMNS = [
    "PN", "Descrição", "UM", "SALDO original ORION", "Falta inicial",
    "M221 encontrado (bruto)", "P277 encontrado (bruto)", "Total encontrado (bruto)",
    "Cobertura potencial (não validada)", "Quantidade confirmada (bruta)",
    "Disponível aplicado ao gap", "Indisponível justificada",
    "Falta física restante", "Falta sem justificativa", "Registros pendentes",
    "Situação do PN", "Próxima ação", "Evidência e linhas STK ORIGINAL",
    "Validações atuais por depósito", "Justificativa atual", "Histórico de revisões",
    "Encaminhamento provisório?",
]
STOCK_WIDTHS = [
    24, 38, 9, 22, 19, 23, 23, 24, 29, 26, 25, 26, 24, 25,
    20, 28, 66, 75, 80, 65, 90, 24,
]

QUALITY_COLUMNS = [
    "PN", "Descrição", "UM", "Falta recebida de Estoque",
    "Falta física após Estoque", "Q239 ativo (bruto)", "Q273 ativo (bruto)",
    "Q273 informativo (não somar)", "Depósito CQ priorizado",
    "Cobertura potencial (não validada)", "Quantidade confirmada (bruta)",
    "Liberado aplicado ao gap", "Reprovado justificado", "Refugo justificado",
    "Falta física após CQ", "Falta sem explicação", "Registros pendentes CQ",
    "Registros pendentes Estoque", "Situação do PN", "Próxima ação",
    "Evidência e linhas STK ORIGINAL", "Validações atuais por depósito",
    "Justificativa atual", "Histórico de revisões", "Encaminhamento provisório?",
]
QUALITY_WIDTHS = [
    24, 38, 9, 28, 25, 21, 21, 29, 26, 29, 27, 25, 26,
    24, 23, 24, 22, 25, 28, 66, 80, 85, 65, 90, 24,
]

STOCK_TABS = (
    ("PNs Críticos", "Todos os PNs críticos do Cenário ORION, inclusive sem registros em M221/P277."),
    ("PNs Encontrados", "Somente PNs com registros válidos no M221 ou P277."),
    ("Aguardando Almoxarifado", "PNs com pelo menos uma evidência M221/P277 ainda sem validação do PCM/Almoxarifado."),
    ("Sem Registros", "PNs críticos não encontrados em M221/P277. Isso não demonstra inexistência em outros depósitos."),
    ("Encaminhados à Qualidade", "PNs cujo gap não explicado foi encaminhado à Qualidade, inclusive encaminhamentos provisórios."),
)

QUALITY_TABS = (
    ("PNs Recebidos", "Todos os PNs encaminhados pela Divergência de Estoque, com ou sem evidência no CQ."),
    ("Encontrados Q239", "Somente PNs com Q239 como depósito ativo, prioritário na análise de CQ."),
    ("Encontrados Q273", "Somente PNs com Q273 como depósito ativo na ausência de Q239; refugo depende de comprovação."),
    ("Aguardando CQ", "PNs com pelo menos uma evidência ativa Q239/Q273 ainda não confirmada pelo CQ."),
    ("Sem Registros nesses depósitos", "PNs encaminhados que não foram encontrados em Q239 nem Q273. Não significa refugo confirmado."),
)


def _decision_label(value):
    return DECISION_LABELS.get(value, value or "Aguardando validação")


def _excel_text(value):
    """Avoid silent truncation at Excel's 32,767 character per-cell limit."""
    value = str(value or "")
    return value if len(value) <= 32000 else value[:31800] + "\n[Texto longo: consulte o histórico no ORION]"


def _source_lines(candidates, *, informative=False):
    rendered = []
    for candidate in candidates:
        rows = ", ".join(str(x) for x in (candidate.get("source_rows") or [])) or "não identificadas"
        descriptor = (
            f"{candidate.get('depot') or '?'} | Centro: {candidate.get('center') or 'não informado'} "
            f"| Quantidade: {candidate.get('quantity') or 0} {candidate.get('unit') or ''} "
            f"| ORIGINAL, linhas {rows}"
        )
        rendered.append(descriptor + (" | INFORMATIVO, NÃO SOMAR" if informative else ""))
    return _excel_text("\n".join(rendered)) if rendered else "STK ORIGINAL: nenhum registro nesses depósitos"


def _review_history(candidate):
    history = candidate.get("validation_history") or []
    if not history:
        return "Sem validações registradas"
    rendered = []
    for version in history:
        reason = version.get("justification") or version.get("evidence") or ""
        rendered.append(
            f"Revisão #{version.get('id') or '?'} | {version.get('created_at') or 'data não informada'} "
            f"| {_decision_label(version.get('decision'))} "
            f"| {_decision_label(version.get('accounting'))} "
            f"| {version.get('confirmed_quantity') or 0} {candidate.get('unit') or ''} "
            f"| {reason or 'Sem justificativa registrada'}"
        )
    return _excel_text("\n".join(rendered))


def _candidate_audit(candidates, *, informative=()):
    """Keep validation/evidence associated with its own depot and SAP rows."""
    current = []
    explanations = []
    versions = []
    confirmed = 0.0
    for candidate in candidates:
        state = candidate.get("validation") or {}
        place = f"{candidate.get('depot') or '?'} / centro {candidate.get('center') or 'não informado'}"
        if state:
            qty = float(state.get("confirmed_quantity") or 0)
            confirmed += qty
            current.append(
                f"{place}: {_decision_label(state.get('decision'))}, "
                f"{_decision_label(state.get('accounting'))}; "
                f"{qty:g} {candidate.get('unit') or ''} confirmado(s), "
                f"registro #{state.get('id') or '?'} em {state.get('created_at') or 'data não informada'}"
            )
            reason = state.get("justification") or state.get("evidence")
            if reason:
                explanations.append((place, str(reason)))
        else:
            current.append(f"{place}: AGUARDANDO CONFIRMAÇÃO")
        versions.append(f"{place}\n{_review_history(candidate)}")
    for candidate in informative:
        place = f"{candidate.get('depot') or '?'} / centro {candidate.get('center') or 'não informado'}"
        current.append(f"{place}: SOMENTE INFORMATIVO; Q239 tem prioridade; não somar nem validar automaticamente")
    if len(explanations) == 1:
        # Keep a user's comment in its own cell, safely escaped by _safe().
        reason_text = explanations[0][1]
    else:
        reason_text = "\n".join(f"{place}: {reason}" for place, reason in explanations)
    return (
        confirmed,
        _excel_text("\n".join(current)) or "Sem registro nos depósitos investigados",
        _excel_text(reason_text) if reason_text else "Sem justificativa registrada",
        _excel_text("\n\n".join(versions)) if versions else "Sem revisões",
    )


def _stock_row(item, *, matched):
    deficit = float(item.get("negative_balance") or 0)
    candidates = item.get("candidates") or []
    raw = sum(float(c.get("quantity") or 0) for c in candidates)
    confirmed, current, reason, history = _candidate_audit(candidates)
    physical = float(item.get("projected_gap_remaining", deficit) or 0) if matched else deficit
    unexplained = float(item.get("unexplained_quantity", deficit) or 0) if matched else deficit
    return [
        item.get("material"), item.get("description"), item.get("unit"),
        item.get("balance", -deficit), deficit,
        sum(float(c.get("quantity") or 0) for c in candidates if c.get("depot") == "M221"),
        sum(float(c.get("quantity") or 0) for c in candidates if c.get("depot") == "P277"),
        raw, min(deficit, raw), confirmed,
        item.get("recovered_quantity_applied", 0),
        item.get("justified_unavailable_applied", 0),
        physical, unexplained, sum(not c.get("validation") for c in candidates),
        _status_label(item.get("workflow_status") if matched else "QUALITY_4_2"),
        item.get("next_step") if matched else "Sem registro M221/P277; encaminhado à Qualidade",
        _source_lines(candidates) if matched else "STK ORIGINAL: nenhum registro M221/P277 para este PN",
        current, reason, history, "Sim" if item.get("routing_provisional") else "Não",
    ]


def _quality_row(item):
    candidates = item.get("candidates") or []
    additional = item.get("additional_q273_evidence") or []
    confirmed, current, reason, history = _candidate_audit(candidates, informative=additional)
    evidence = _source_lines(candidates)
    if additional:
        evidence += "\n" + _source_lines(additional, informative=True)
    return [
        item.get("material"), item.get("description"), item.get("unit"),
        item.get("gap_from_41", 0), item.get("physical_gap_from_41", 0),
        sum(float(c.get("quantity") or 0) for c in candidates if c.get("depot") == "Q239"),
        sum(float(c.get("quantity") or 0) for c in candidates if c.get("depot") == "Q273"),
        sum(float(c.get("quantity") or 0) for c in additional),
        item.get("active_depot") or "Sem Q239/Q273",
        item.get("theoretical_coverage_limit", 0),
        confirmed, item.get("quality_released_applied", 0),
        item.get("quality_rejected_applied", 0), item.get("quality_scrap_applied", 0),
        item.get("physical_gap_after_quality", 0), item.get("unexplained_after_quality", 0),
        sum(not c.get("validation") for c in candidates),
        item.get("pending_stock_reviews", 0),
        _status_label(item.get("workflow_status")), item.get("next_step"),
        _excel_text(evidence), current, reason, history,
        "Sim" if item.get("routing_provisional") else "Não",
    ]


def _sorted(items, *, deficit):
    return sorted(items, key=lambda item: (-float(item.get(deficit) or 0), str(item.get("material") or "")))


def _export_stock(book, report, *, subtitle):
    found = {item["material"]: item for item in report.get("materials", [])}
    unmatched = {item["material"]: item for item in report.get("unmatched_materials", [])}
    all_items = {**found, **unmatched}
    forwarded = report.get("workflow", {}).get("forward_to_quality") or []
    subsets = (
        list(all_items.values()),
        list(found.values()),
        [x for x in found.values() if any(not c.get("validation") for c in x.get("candidates") or [])],
        list(unmatched.values()),
        [all_items[entry["material"]] for entry in forwarded if entry["material"] in all_items],
    )
    expected = report.get("summary", {})
    counts = (
        expected.get("critical_materials"),
        expected.get("materials_with_candidates"),
        # The web uses unique PNs with any unreviewed candidate, not just workflow status.
        None,
        expected.get("materials_without_candidates"),
        expected.get("items_to_quality"),
    )
    for (name, description), selection, count in zip(STOCK_TABS, subsets, counts):
        records = [_stock_row(item, matched=item["material"] in found) for item in _sorted(selection, deficit="negative_balance")]
        if count is not None and len(records) != count:
            raise ValueError(f"Indicador de Estoque inconsistente: {name}.")
        _rows(
            book.create_sheet(name), STOCK_COLUMNS, STOCK_WIDTHS, records,
            title=name.upper(),
            subtitle=f"Indicador do painel: {len(records)} PN(s) | {subtitle}",
            note=description + " Saldos encontrados não são liberações e nenhum SALDO oficial é alterado.",
            status_column=16,
        )


def _export_quality(book, report, *, subtitle):
    materials = report.get("materials", [])
    subsets = (
        materials,
        [x for x in materials if x.get("active_depot") == "Q239"],
        [x for x in materials if x.get("active_depot") == "Q273"],
        [x for x in materials if any(not c.get("validation") for c in x.get("candidates") or [])],
        [x for x in materials if x.get("active_depot") is None],
    )
    counts = report.get("summary", {})
    expected = (
        counts.get("items_received_from_41"),
        counts.get("materials_with_q239"),
        counts.get("materials_with_q273"),
        None,
        counts.get("materials_without_quality_records"),
    )
    for (name, description), selection, count in zip(QUALITY_TABS, subsets, expected):
        records = [_quality_row(item) for item in _sorted(selection, deficit="gap_from_41")]
        if count is not None and len(records) != count:
            raise ValueError(f"Indicador de Qualidade inconsistente: {name}.")
        _rows(
            book.create_sheet(name), QUALITY_COLUMNS, QUALITY_WIDTHS, records,
            title=name.upper(),
            subtitle=f"Indicador do painel: {len(records)} PN(s) | {subtitle}",
            note=description + " Quantidades por PN e UM; não somar refugo informativo ou confundir evidência com liberação.",
            status_column=19,
        )


def export_investigation_excel(
    kind: str, *, report: dict, scenario: dict,
    stock_report: dict | None = None, quality_report: dict | None = None,
) -> tuple[bytes, str, str]:
    """Export the five actual dashboard KPI cohorts, without unrelated tabs."""
    if kind not in {"stock", "quality"}:
        raise ValueError("Tipo de investigação inválido.")
    book = Workbook()
    book.remove(book.active)
    source = report.get("source") or {}
    month = report.get("reference_month") or scenario.get("reference_month") or "SEM_MES"
    stock_name = source.get("filename") or "STK SAP"
    subtitle = f"{month} · {stock_name} · aba ORIGINAL · cenário {report.get('scenario_id') or 'não informado'}"
    if kind == "stock":
        _export_stock(book, stock_report if stock_report is not None else report, subtitle=subtitle)
    else:
        _export_quality(book, quality_report if quality_report is not None else report, subtitle=subtitle)
    label = "Qualidade" if kind == "quality" else "Estoque"
    stamp = str(month).replace("-", "_")
    output = BytesIO()
    book.save(output)
    return output.getvalue(), f"ORION_{label}_Investigacao_{stamp}.xlsx", XLSX_MIME
