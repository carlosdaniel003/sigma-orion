"""Approval gate and authoritative Excel handoff after both DPP investigations."""
from hashlib import sha256
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services import dpp_scenario_service as scenarios
from app.services.dpp_canonical_export_service import export_monthly_scenario_excel
from app.services.dpp_export_formula_fidelity_service import finalize_and_audit_dpp_formulas
from test_dpp_canonical_export import _template_bytes
from test_dpp_stock_divergence import MIME, stock_bytes


def _stock():
    return stock_bytes([
        ["PN-A", "Stock", 8, "M221", 0, 0, "UN", "1063"],
        ["PN-A", "CQ", 5, "Q239", 0, 0, "UN", "1063"],
    ])


def _scenario(stock):
    scenarios._SCENARIOS.clear()
    return scenarios.register_monthly_scenario(
        materials=[
            {
                "material": "PN-A", "description": "Material A", "um": "UN",
                "group_origin": "LOCAL", "consumption_by_model": {"MODELO A": 1.0},
                "stock_sap_effective": 0.0, "explosion": 0.0, "stock_op": 0.0,
                "price": 2,
            },
            {
                "material": "PN-B", "description": "Material B", "um": "UN",
                "group_origin": "LOCAL", "consumption_by_model": {"MODELO A": 0.4},
                "stock_sap_effective": 0.0, "explosion": 0.0, "stock_op": 0.0,
                "price": 3,
            },
            {
                "material": "PN-C", "description": "Material C", "um": "UN",
                "group_origin": "LOCAL", "consumption_by_model": {"MODELO A": 0.0},
                "stock_sap_effective": 7.0, "explosion": 0.0, "stock_op": 0.0,
                "price": 1,
            },
        ],
        models=[
            {"name": "MODELO A", "kit_pgd": 10.0},
            {"name": "MODELO B", "kit_pgd": 0.0},
        ],
        reference_month="2026-07",
        base_summary={}, scope="test", capabilities={},
        sources=[{"id": "stock", "filename": "STK JULHO.xlsx",
                  "fingerprint_sha256": sha256(stock).hexdigest()}],
        pending=[], diagnostics={}, pgd_mapping={},
    )


def _post(client, url, sid, stock, fields=None):
    return client.post(
        url,
        data={"scenario_id": sid, **(fields or {})},
        files={"stock": ("STK JULHO.xlsx", BytesIO(stock), MIME)},
    )


def _record_stock(client, sid, stock, *, qty=4):
    report = _post(client, "/api/dpp/investigations/stock-divergence", sid, stock).json()
    identifier = report["materials"][0]["candidates"][0]["candidate_id"]
    saved = _post(client, "/api/dpp/investigations/stock-divergence/review", sid, stock, {
        "candidate_id": identifier, "decision": "AVAILABLE",
        "accounting": "NOT_INCLUDED", "confirmed_quantity": str(qty),
    })
    assert saved.status_code == 200, saved.text
    return identifier


def _record_quality(client, sid, stock, *, qty=2):
    report = _post(client, "/api/dpp/investigations/quality", sid, stock)
    assert report.status_code == 200, report.text
    identifier = report.json()["materials"][0]["candidates"][0]["candidate_id"]
    saved = _post(client, "/api/dpp/investigations/quality/review", sid, stock, {
        "candidate_id": identifier, "decision": "RELEASED",
        "accounting": "NOT_INCLUDED", "confirmed_quantity": str(qty),
    })
    assert saved.status_code == 200, saved.text
    return identifier


def test_full_month_application_requires_human_confirmations_and_two_phase_preview(tmp_path, monkeypatch):
    monkeypatch.setattr(scenarios, "SCENARIO_STATE_PATH", tmp_path / "applied.json")
    stock = _stock()
    original = _scenario(stock)
    sid = original["scenario_id"]
    assert original["summary"]["materials_to_investigate"] == 2
    assert {m["material"]: m["balance"] for m in original["materials"]}["PN-A"] == -10

    try:
        with TestClient(app) as client:
            first = _post(client, "/api/dpp/monthly/gap/prepare", sid, stock)
            assert first.status_code == 200, first.text
            assert first.json()["ready"] is False
            assert first.json()["stock_pending"] == 1
            assert _post(client, "/api/dpp/monthly/gap/apply", sid, stock, {
                "approval_token": first.json()["approval_token"], "confirm": "true",
            }).status_code == 400

            _record_stock(client, sid, stock)
            mid = _post(client, "/api/dpp/monthly/gap/prepare", sid, stock).json()
            assert mid["ready"] is False
            assert mid["stock_pending"] == 0
            assert mid["quality_pending"] == 1
            _record_quality(client, sid, stock)

            preview = _post(client, "/api/dpp/monthly/gap/prepare", sid, stock).json()
            assert preview["ready"] is True
            assert preview["stock_pending"] == preview["quality_pending"] == 0
            assert preview["materials_total"] == 3
            assert preview["unresolved_materials"] == 2
            assert preview["total_stock_confirmed"] == 6
            assert preview["by_unit"][0]["unit"] == "UN"
            assert preview["by_unit"][0]["from_stock"] == 4
            assert preview["by_unit"][0]["from_quality"] == 2

            fields = {"approval_token": preview["approval_token"]}
            assert _post(client, "/api/dpp/monthly/gap/apply", sid, stock, fields).status_code == 400
            assert _post(client, "/api/dpp/monthly/gap/apply", sid, stock,
                         {**fields, "confirm": "true"}).status_code == 400

            applied_response = _post(client, "/api/dpp/monthly/gap/apply", sid, stock, {
                **fields, "confirm": "true", "accept_unresolved": "true",
            })
            assert applied_response.status_code == 200, applied_response.text
            payload = applied_response.json()
            new_scenario = payload["scenario"]
            new_sid = payload["applied_scenario_id"]
            assert new_sid != sid
            assert payload["original_scenario_id"] == sid
            assert new_scenario["gap_approval"]["stock_sha256"] == sha256(stock).hexdigest()
            assert new_scenario["gap_approval"]["pending_validations_at_approval"] == 0
            assert new_scenario["status"] == "GAP_APLICADO"
            before_by_id = {x["material"]: x for x in scenarios.get_monthly_scenario(sid)["materials"]}
            after_by_id = {x["material"]: x for x in new_scenario["materials"]}
            assert before_by_id["PN-A"]["balance"] == -10
            assert after_by_id["PN-A"]["approved_gap_added_stock"] == 6
            assert after_by_id["PN-A"]["stock_sap_effective"] == 0
            assert after_by_id["PN-A"]["stock_total"] == 6
            assert after_by_id["PN-A"]["balance"] == -4
            assert after_by_id["PN-A"]["amount"] == -8
            assert after_by_id["PN-B"]["balance"] == -4
            assert after_by_id["PN-C"]["balance"] == 7
            assert new_scenario["summary"]["materials_to_investigate"] == 2

            assert _post(client, "/api/dpp/monthly/gap/apply", sid, stock, {
                **fields, "confirm": "true", "accept_unresolved": "true",
            }).status_code == 400
            assert _post(client, "/api/dpp/monthly/gap/prepare", new_sid, stock).status_code == 400

            with pytest.raises(ValueError, match="versão fechada"):
                scenarios.recalculate_monthly_scenario(new_sid, {"MODELO A": 12})

            assert scenarios.get_latest_monthly_scenario()["scenario_id"] == new_sid

            # The canonical monthly workbook must match *approved* SALDO while
            # preserving the raw SAP column and every downstream formula.
            content, filename, _ = export_monthly_scenario_excel(
                scenario_id=new_sid, template_content=_template_bytes(),
                template_filename="DPP JUN_2026.xlsx",
            )
            assert filename == "DPP_ORION_2026_07.xlsx"
            book = load_workbook(BytesIO(content), data_only=False)
            try:
                assert book["ORION_CONFIRMACOES"].sheet_state == "hidden"
                approved_sheet = book["ORION_CONFIRMACOES"]
                assert approved_sheet["A2"].value == "PN-A"
                assert approved_sheet["B2"].value == 4
                assert approved_sheet["C2"].value == 2
                assert approved_sheet["D2"].value == "=B2+C2"
                assert approved_sheet["E2"].value == -10

                dpp = book["DPP"]
                rows = {dpp.cell(row, 1).value: row for row in range(5, 8)}
                a = rows["PN-A"]
                b = rows["PN-B"]
                assert dpp.cell(a, 10).value == 0  # STK SAP unchanged
                assert "ORION_CONFIRMACOES" in dpp.cell(a, 14).value
                assert "VLOOKUP" in dpp.cell(a, 14).value
                assert dpp.cell(a, 15).value == f"=N{a}-I{a}"
                assert dpp.cell(a, 17).value == f"=P{a}*O{a}"
                assert "ORION_CONFIRMACOES" not in dpp.cell(b, 14).value
            finally:
                book.close()

            audited = finalize_and_audit_dpp_formulas(content, new_scenario)
            assert audited

            scenarios._SCENARIOS.clear()
            assert scenarios.restore_persisted_monthly_scenario() is True
            assert scenarios.get_latest_monthly_scenario()["scenario_id"] == new_sid
            restored_original = scenarios.get_monthly_scenario(sid)
            assert restored_original is not None
            assert {r["material"]: r for r in restored_original["materials"]}["PN-A"]["balance"] == -10
            assert {r["material"]: r for r in scenarios.get_monthly_scenario(new_sid)["materials"]}["PN-A"]["balance"] == -4
    finally:
        scenarios._SCENARIOS.clear()


def test_stale_confirmation_token_rejects_changed_validation(tmp_path, monkeypatch):
    monkeypatch.setattr(scenarios, "SCENARIO_STATE_PATH", tmp_path / "stale.json")
    stock = _stock()
    scenario = _scenario(stock)
    sid = scenario["scenario_id"]
    try:
        with TestClient(app) as client:
            identifier = _record_stock(client, sid, stock)
            _record_quality(client, sid, stock)
            previous = _post(client, "/api/dpp/monthly/gap/prepare", sid, stock).json()
            assert previous["ready"] is True
            revised = _post(client, "/api/dpp/investigations/stock-divergence/review", sid, stock, {
                "candidate_id": identifier, "decision": "AVAILABLE",
                "accounting": "NOT_INCLUDED", "confirmed_quantity": "3",
            })
            assert revised.status_code == 200, revised.text
            outdated = _post(client, "/api/dpp/monthly/gap/apply", sid, stock, {
                "approval_token": previous["approval_token"],
                "confirm": "true", "accept_unresolved": "true",
            })
            assert outdated.status_code == 400
            assert "mudaram" in outdated.json()["detail"]
            assert scenarios.get_latest_monthly_scenario()["scenario_id"] == sid
    finally:
        scenarios._SCENARIOS.clear()
