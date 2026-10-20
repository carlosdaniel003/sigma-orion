"""Audit regression: cap attribution per PN, preserve evidence, flag pending checks."""
from uuid import uuid4
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import routes
from app.main import app
from test_dpp_stock_divergence import MIME, stock_bytes


def scenario(pns):
    return {
        "scenario_id": uuid4().hex,
        "reference_month": "2026-07",
        "sources": [
            {"id": "base_dpp", "filename": "DPP JUN_2026.06.02.xlsx"},
            {"id": "stock", "filename": "STK JULHO.xlsx"},
        ],
        "materials": [
            {"material": pn, "description": "Teste " + pn, "um": "UN", "balance": -gap, "status": "INVESTIGAR"}
            for pn, gap in pns
        ],
    }


def line(pn, qty, depot, center="1063"):
    return [pn, "Teste " + pn, qty, depot, 0, 0, "UN", center]


def post(client, path, current, stock, more=None):
    return client.post(path, data={"scenario_id": current["scenario_id"], **(more or {})},
                       files={"stock": ("STK JULHO.xlsx", stock, MIME)})


def worksheet(client, path, current, stock):
    result = post(client, path, current, stock)
    assert result.status_code == 200, result.text
    return load_workbook(BytesIO(result.content), data_only=True)


def test_stock_applies_only_deficit_and_preserves_raw_excess(monkeypatch):
    current = scenario([("PN-1", 5), ("PN-SEM", 2)])
    stock = stock_bytes([line("PN-1", 6, "M221"), line("PN-1", 8, "P277")])
    monkeypatch.setattr(routes, "get_monthly_scenario", lambda s: current if s == current["scenario_id"] else None)

    with TestClient(app) as client:
        result = post(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        first, second = [c["candidate_id"] for c in result["materials"][0]["candidates"]]
        for identifier, decision, qty, justification, accounting in (
            (first, "AVAILABLE", 6, "", "NOT_INCLUDED"),
            (second, "UNAVAILABLE", 8, "Oito peças não localizadas.", "UNKNOWN"),
        ):
            saved = post(client, "/api/dpp/investigations/stock-divergence/review", current, stock, {
                "candidate_id": identifier, "decision": decision,
                "accounting": accounting, "confirmed_quantity": str(qty), "evidence": justification,
            })
            assert saved.status_code == 200, saved.text

        data = saved.json()
        item = data["materials"][0]
        assert item["confirmed_increment_outside_dpp"] == 6
        assert item["recovered_quantity_applied"] == 5
        assert item["justified_unavailable_quantity"] == 8
        assert item["justified_unavailable_applied"] == 0
        assert item["confirmed_excess_not_attributed"] == 9
        assert item["projected_gap_remaining"] == 0
        assert item["unexplained_quantity"] == 0
        assert item["theoretical_coverage_limit"] == 5
        assert item["candidates"][0]["recovered_applied"] == 5
        assert item["candidates"][1]["justified_unavailable_applied"] == 0
        assert data["summary"]["items_to_quality"] == 1  # PN-SEM has no special depot.

        book = worksheet(client, "/api/dpp/investigations/stock-divergence/export", current, stock)
        materials = {row[0]: row for row in book["PNs Críticos"].iter_rows(min_row=6, values_only=True) if row[0]}
        audited = materials["PN-1"]
        assert audited[3] == -5
        assert audited[7] == 14
        assert audited[8] == 5
        assert audited[9] == 14  # two independent confirmations, 6 + 8
        assert audited[10] == 5  # available applied only up to deficit
        assert audited[11] == 0
        assert audited[15] == "Finalizado"
        assert materials["PN-SEM"][13] == 2
        assert "M221" in audited[17] and "P277" in audited[17]
        assert "6 UN" in audited[18] and "8 UN" in audited[18]
        # Per-unit deficit and potential coverage can be audited in the complete cohort.
        assert sum(row[4] for row in materials.values()) == 7
        assert sum(row[8] for row in materials.values()) == 5
        assert len([row for row in book["Encaminhados à Qualidade"].iter_rows(min_row=6, values_only=True) if row[0]]) == 1
        book.close()


def test_pending_stock_reviews_create_only_provisional_quality_handoff(monkeypatch):
    current = scenario([("PN-1", 8)])
    stock = stock_bytes([line("PN-1", 3, "M221"), line("PN-1", 4, "P277"), line("PN-1", 12, "Q239")])
    monkeypatch.setattr(routes, "get_monthly_scenario", lambda s: current if s == current["scenario_id"] else None)

    with TestClient(app) as client:
        first = post(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        first_id = first["materials"][0]["candidates"][0]["candidate_id"]
        applied = post(client, "/api/dpp/investigations/stock-divergence/review", current, stock, {
            "candidate_id": first_id, "decision": "AVAILABLE",
            "accounting": "NOT_INCLUDED", "confirmed_quantity": "3",
        })
        assert applied.status_code == 200
        entry = applied.json()["workflow"]["forward_to_quality"][0]
        assert entry["recovered_quantity"] == 3
        assert entry["routing_provisional"] is True
        assert entry["pending_reviews"] == 1
        assert applied.json()["summary"]["provisional_handoffs"] == 1

        quality = post(client, "/api/dpp/investigations/quality", current, stock)
        assert quality.status_code == 200
        candidate_id = quality.json()["materials"][0]["candidates"][0]["candidate_id"]
        saved = post(client, "/api/dpp/investigations/quality/review", current, stock, {
            "candidate_id": candidate_id, "decision": "RELEASED",
            "accounting": "NOT_INCLUDED", "confirmed_quantity": "8",
        })
        assert saved.status_code == 200, saved.text
        data = saved.json()
        quality_item = data["materials"][0]
        assert quality_item["quality_released"] == 8
        assert quality_item["quality_released_applied"] == 5
        assert quality_item["confirmed_excess_not_attributed"] == 3
        assert quality_item["physical_gap_after_quality"] == 0
        assert quality_item["unexplained_after_quality"] == 0
        assert quality_item["routing_provisional"] is True
        assert "provisória" in quality_item["next_step"]

        book = worksheet(client, "/api/dpp/investigations/quality/export", current, stock)
        item = next(row for row in book["PNs Recebidos"].iter_rows(min_row=6, values_only=True) if row[0] == "PN-1")
        assert item[11] == 5  # released applied, not eight confirmed
        assert item[9] == 5   # potential bounded by incoming gap
        assert item[24] == "Sim"
        assert item[10] == 8  # raw confirmation is still visible
        assert "Q239" in item[20] and "8 UN" in item[21]
        book.close()


def test_quality_rejection_explains_only_remaining_missing_amount(monkeypatch):
    current = scenario([("PN-1", 3)])
    stock = stock_bytes([line("PN-1", 5, "Q239"), line("PN-1", 11, "Q273")])
    monkeypatch.setattr(routes, "get_monthly_scenario", lambda s: current if s == current["scenario_id"] else None)
    with TestClient(app) as client:
        previous = post(client, "/api/dpp/investigations/quality", current, stock).json()
        identifier = previous["materials"][0]["candidates"][0]["candidate_id"]
        validated = post(client, "/api/dpp/investigations/quality/review", current, stock, {
            "candidate_id": identifier, "decision": "REJECTED",
            "accounting": "UNKNOWN", "confirmed_quantity": "5",
            "justification": "Cinco peças reprovadas pelo CQ.",
        })
        assert validated.status_code == 200
        item = validated.json()["materials"][0]
        assert item["quality_rejected_justified"] == 5
        assert item["quality_rejected_applied"] == 3
        assert item["confirmed_excess_not_attributed"] == 2
        assert item["quality_released_applied"] == 0
        assert item["physical_gap_after_quality"] == 3  # Rejection never recovers stock.
        assert item["unexplained_after_quality"] == 0
        assert item["workflow_status"] == "FINISHED"
        book = worksheet(client, "/api/dpp/investigations/quality/export", current, stock)
        report_row = next(row for row in book["Encontrados Q239"].iter_rows(min_row=6, values_only=True) if row[0] == "PN-1")
        assert report_row[12] == 3  # rejection applied, not all five raw
        assert report_row[15] == 0  # unexplained shortage cleared
        assert report_row[10] == 5  # confirmed raw quantity
        assert report_row[7] == 11  # Q273 informative only
        assert "Q273" in report_row[20] and "INFORMATIVO" in report_row[20]
        assert "PN-1" not in [row[0] for row in book["Encontrados Q273"].iter_rows(min_row=6, values_only=True)]
        book.close()
