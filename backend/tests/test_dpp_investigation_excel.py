"""Readable investigation spreadsheet downloads: contents, traceability and safety."""
from io import BytesIO
from hashlib import sha256
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import routes
from app.main import app
from test_dpp_stock_divergence import MIME, stock_bytes


def scenario():
    return {
        "scenario_id": uuid4().hex,
        "reference_month": "2026-07",
        "sources": [
            {"id": "base_dpp", "filename": "DPP JUN_2026.06.02.xlsx"},
            {"id": "stock", "filename": "STK JULHO.xlsx"},
            {"id": "wiu", "filename": "WIU JULHO.xlsx"},
            {"id": "explosion", "filename": "Explosão JULHO.xlsm"},
            {"id": "pgd", "filename": "PGD BASE.xlsx"},
        ],
        "materials": [
            {"material": "PN-CQ", "description": "Componente em CQ", "um": "UN", "balance": -10, "status": "INVESTIGAR"},
            {"material": "PN-V", "description": "Material sem registro especial", "um": "UN", "balance": -4, "status": "INVESTIGAR"},
            {"material": "PN-OK", "description": "Material normal", "um": "UN", "balance": 100, "status": "OK"},
        ],
    }


def rows():
    return [
        ["PN-CQ", "Componente em CQ", 12, "M221", 0, 0, "UN", "1063"],
        ["PN-CQ", "Componente em CQ", 7, "Q239", 0, 0, "UN", "1063"],
        ["PN-CQ", "Componente em CQ", 3, "Q273", 0, 0, "UN", "1063"],
    ]


def upload(client, path, current, stock):
    return client.post(
        path,
        data={"scenario_id": current["scenario_id"]},
        files={"stock": ("STK JULHO.xlsx", stock, MIME)},
    )


def spreadsheet(response):
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"PK")
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment;" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content), data_only=False)


STOCK_TABS = [
    "PNs Críticos", "PNs Encontrados", "Aguardando Almoxarifado",
    "Sem Registros", "Encaminhados à Qualidade",
]
QUALITY_TABS = [
    "PNs Recebidos", "Encontrados Q239", "Encontrados Q273",
    "Aguardando CQ", "Sem Registros nesses depósitos",
]


def pn_rows(book, tab):
    return {row[0]: row for row in book[tab].iter_rows(min_row=6, values_only=True) if row[0] and row[0] != "Nenhum registro para os critérios desta investigação."}


def test_stock_excel_tabs_match_each_dashboard_indicator(monkeypatch):
    current = scenario()
    stock = stock_bytes(rows())
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        dashboard = upload(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        book = spreadsheet(upload(client, "/api/dpp/investigations/stock-divergence/export", current, stock))
        assert book.sheetnames == STOCK_TABS
        summary = dashboard["summary"]
        assert len(pn_rows(book, "PNs Críticos")) == summary["critical_materials"] == 2
        assert len(pn_rows(book, "PNs Encontrados")) == summary["materials_with_candidates"] == 1
        assert len(pn_rows(book, "Aguardando Almoxarifado")) == 1
        assert len(pn_rows(book, "Sem Registros")) == summary["materials_without_candidates"] == 1
        assert len(pn_rows(book, "Encaminhados à Qualidade")) == summary["items_to_quality"] == 1
        assert set(pn_rows(book, "PNs Críticos")) == {"PN-CQ", "PN-V"}
        assert set(pn_rows(book, "PNs Encontrados")) == {"PN-CQ"}
        assert set(pn_rows(book, "Aguardando Almoxarifado")) == {"PN-CQ"}
        assert set(pn_rows(book, "Sem Registros")) == {"PN-V"}
        assert set(pn_rows(book, "Encaminhados à Qualidade")) == {"PN-V"}
        pn = pn_rows(book, "PNs Encontrados")["PN-CQ"]
        assert pn[3] == -10
        assert pn[5] == 12  # M221 raw stock, not confirmed release
        assert pn[6] == 0
        assert pn[7] == 12
        assert pn[8] == 10  # coverage bounded by individual deficit
        assert pn[9] == 0   # no human confirmation
        assert pn[10] == 0  # no stock reduction
        assert pn[13] == 10
        assert pn[14] == 1
        assert "ORIGINAL, linhas 2" in pn[17] and "1063" in pn[17]
        assert "AGUARDANDO CONFIRMAÇÃO" in pn[18]
        assert "STK JULHO.xlsx" in book["PNs Críticos"]["A2"].value
        assert all(s.freeze_panes == "C6" and s.auto_filter.ref for s in book.worksheets)
        book.close()


def test_quality_excel_tabs_match_panel_and_q239_priority(monkeypatch):
    current = scenario()
    stock = stock_bytes(rows())
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        book = spreadsheet(upload(client, "/api/dpp/investigations/quality/export", current, stock))
        assert book.sheetnames == QUALITY_TABS
        assert set(pn_rows(book, "PNs Recebidos")) == {"PN-V"}
        assert not pn_rows(book, "Encontrados Q239")
        assert not pn_rows(book, "Encontrados Q273")
        assert not pn_rows(book, "Aguardando CQ")
        assert set(pn_rows(book, "Sem Registros nesses depósitos")) == {"PN-V"}
        assert "não" in book["Sem Registros nesses depósitos"]["A3"].value.lower()
        book.close()

        first = upload(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        candidate = first["materials"][0]["candidates"][0]["candidate_id"]
        saved = client.post(
            "/api/dpp/investigations/stock-divergence/review",
            data={"scenario_id": current["scenario_id"], "candidate_id": candidate,
                  "decision": "UNAVAILABLE", "accounting": "UNKNOWN",
                  "confirmed_quantity": "1", "evidence": "Uma peça indisponível."},
            files={"stock": ("STK JULHO.xlsx", stock, MIME)},
        )
        assert saved.status_code == 200
        quality = upload(client, "/api/dpp/investigations/quality", current, stock).json()
        book = spreadsheet(upload(client, "/api/dpp/investigations/quality/export", current, stock))
        assert len(pn_rows(book, "PNs Recebidos")) == quality["summary"]["items_received_from_41"] == 2
        assert len(pn_rows(book, "Encontrados Q239")) == quality["summary"]["materials_with_q239"] == 1
        assert len(pn_rows(book, "Encontrados Q273")) == quality["summary"]["materials_with_q273"] == 0
        assert len(pn_rows(book, "Aguardando CQ")) == 1
        assert len(pn_rows(book, "Sem Registros nesses depósitos")) == quality["summary"]["materials_without_quality_records"] == 1
        cq = pn_rows(book, "Encontrados Q239")["PN-CQ"]
        assert cq[3] == 9  # one unit already justified in Stock
        assert cq[5] == 7  # active Q239
        assert cq[6] == 0  # Q273 is not active when Q239 exists
        assert cq[7] == 3  # informational Q273 remains visible
        assert cq[8] == "Q239"
        assert cq[11] == 0  # no confirmed release yet
        assert cq[16] == 1
        assert "ORIGINAL, linhas 3" in cq[20]
        assert "ORIGINAL, linhas 4" in cq[20]
        assert "INFORMATIVO, NÃO SOMAR" in cq[20]
        book.close()


def test_q273_becomes_own_tab_only_when_active(monkeypatch):
    current = scenario()
    stock = stock_bytes([["PN-V", "Refugo a verificar", 3, "Q273", 0, 0, "UN", "1063"]])
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        report = upload(client, "/api/dpp/investigations/quality", current, stock).json()
        book = spreadsheet(upload(client, "/api/dpp/investigations/quality/export", current, stock))
        assert len(pn_rows(book, "PNs Recebidos")) == report["summary"]["items_received_from_41"] == 2
        assert not pn_rows(book, "Encontrados Q239")
        assert set(pn_rows(book, "Encontrados Q273")) == {"PN-V"}
        assert set(pn_rows(book, "Aguardando CQ")) == {"PN-V"}
        assert pn_rows(book, "Encontrados Q273")["PN-V"][6] == 3
        assert pn_rows(book, "Encontrados Q273")["PN-V"][13] == 0  # no scrap confirmed
        book.close()


def test_export_preserves_latest_justification_and_full_review_history(monkeypatch):
    current = scenario()
    stock = stock_bytes(rows())
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        first = upload(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        identifier = first["materials"][0]["candidates"][0]["candidate_id"]
        for text in ("Primeira validação do M221.", "Revisão: lote permanece indisponível."):
            saved = client.post(
                "/api/dpp/investigations/stock-divergence/review",
                data={"scenario_id": current["scenario_id"], "candidate_id": identifier,
                      "decision": "UNAVAILABLE", "accounting": "UNKNOWN",
                      "confirmed_quantity": "3", "evidence": text},
                files={"stock": ("STK JULHO.xlsx", stock, MIME)},
            )
            assert saved.status_code == 200
        book = spreadsheet(upload(client, "/api/dpp/investigations/stock-divergence/export", current, stock))
        reviewed = pn_rows(book, "PNs Encontrados")["PN-CQ"]
        assert reviewed[9] == 3
        assert reviewed[11] == 3
        assert reviewed[19] == "Revisão: lote permanece indisponível."
        assert "Primeira validação do M221." in reviewed[20]
        assert "Revisão: lote permanece indisponível." in reviewed[20]
        assert "PN-CQ" not in pn_rows(book, "Aguardando Almoxarifado")
        assert "PN-CQ" in pn_rows(book, "Encaminhados à Qualidade")
        book.close()


def test_stock_partial_reviews_stay_in_pending_tab_after_handoff(monkeypatch):
    current = scenario()
    stock = stock_bytes(rows() + [["PN-CQ", "Componente em CQ", 4, "P277", 0, 0, "UN", "1063"]])
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        first = upload(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        identifier = next(x["candidate_id"] for x in first["materials"][0]["candidates"] if x["depot"] == "M221")
        saved = client.post(
            "/api/dpp/investigations/stock-divergence/review",
            data={"scenario_id": current["scenario_id"], "candidate_id": identifier,
                  "decision": "AVAILABLE", "accounting": "NOT_INCLUDED",
                  "confirmed_quantity": "2"},
            files={"stock": ("STK JULHO.xlsx", stock, MIME)},
        )
        assert saved.status_code == 200
        book = spreadsheet(upload(client, "/api/dpp/investigations/stock-divergence/export", current, stock))
        assert "PN-CQ" in pn_rows(book, "Aguardando Almoxarifado")
        assert "PN-CQ" in pn_rows(book, "Encaminhados à Qualidade")
        stock_pn = pn_rows(book, "Aguardando Almoxarifado")["PN-CQ"]
        assert stock_pn[10] == 2
        assert stock_pn[14] == 1
        assert stock_pn[21] == "Sim"
        book.close()


def test_export_rejects_wrong_snapshot_and_missing_scenario(monkeypatch):
    current = scenario()
    original = stock_bytes(rows())
    current["sources"][1]["fingerprint_sha256"] = sha256(original).hexdigest()
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        changed = stock_bytes(rows() + [["PN-V", "outro", 5, "M221", 0, 0, "UN", "1063"]])
        for path in (
            "/api/dpp/investigations/stock-divergence/export",
            "/api/dpp/investigations/quality/export",
        ):
            denied = upload(client, path, current, changed)
            assert denied.status_code == 400
            assert "STK" in denied.json()["detail"]
        missing = client.post(
            "/api/dpp/investigations/stock-divergence/export",
            data={"scenario_id": "not-a-scenario"},
            files={"stock": ("STK JULHO.xlsx", original, MIME)},
        )
        assert missing.status_code == 404


def test_export_escapes_untrusted_formula_like_justification(monkeypatch):
    current = scenario()
    stock = stock_bytes(rows())
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    with TestClient(app) as client:
        report = upload(client, "/api/dpp/investigations/stock-divergence", current, stock).json()
        identifier = report["materials"][0]["candidates"][0]["candidate_id"]
        saved = client.post(
            "/api/dpp/investigations/stock-divergence/review",
            data={"scenario_id": current["scenario_id"], "candidate_id": identifier,
                  "decision": "UNAVAILABLE", "accounting": "UNKNOWN",
                  "confirmed_quantity": "2", "evidence": "=HYPERLINK(\"http://example.invalid\", \"unsafe\")"},
            files={"stock": ("STK JULHO.xlsx", stock, MIME)},
        )
        assert saved.status_code == 200
        book = spreadsheet(upload(client, "/api/dpp/investigations/stock-divergence/export", current, stock))
        cell = next(row for row in book["PNs Encontrados"].iter_rows(min_row=6)
                    if row[0].value == "PN-CQ")[19]
        assert cell.data_type != "f"
        assert cell.value.startswith("'=")
        book.close()


def test_generated_excel_has_only_human_operational_labels(monkeypatch):
    import re

    current = scenario()
    stock = stock_bytes(rows())
    monkeypatch.setattr(routes, "get_monthly_scenario",
                        lambda key: current if key == current["scenario_id"] else None)
    forbidden = re.compile(r"fluxograma|\b4\.[123]\b|QUALITY_4_2|SCRAP_4_3", re.IGNORECASE)
    with TestClient(app) as client:
        for route in (
            "/api/dpp/investigations/stock-divergence/export",
            "/api/dpp/investigations/quality/export",
        ):
            book = spreadsheet(upload(client, route, current, stock))
            try:
                for sheet in book.worksheets:
                    for line in sheet.iter_rows(values_only=True):
                        for value in line:
                            if isinstance(value, str):
                                assert not forbidden.search(value), (sheet.title, value)
            finally:
                book.close()
